"""纯内存回归测试：不启动 lifespan，不读取或修改项目数据库，不调用外部平台。"""
import os
os.environ.update(DATABASE_URL="sqlite://", ENVIRONMENT="test", DEMO_SEEDING_ENABLED="false")

import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, func, select, text, inspect
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app
from app.core.config import settings
from app.db.base import Base
from app.core.deps import get_current_user, get_db
from app.core.security import hash_password
from app.models.driver import Driver
from app.models.media_asset import MediaAsset
from app.models.fuel import FuelBalance, FuelCardLookup, FuelRecord, FuelSyncLog
from app.models.nav_preset import NavPreset
from app.models.repair_record import RepairRecord, RepairSettlement, RepairSettlementLine
from app.models.trip_request import TripRequest
from app.models.user import User
from app.models.vehicle import Vehicle
from app.models.workshop import Workshop
from app.services.kunlun_balance_service import import_balance_rows
from app.services.fuel_analysis_service import build_fuel_analysis
from app.services.settlement_recognition_service import recognize_and_save
from app.services.workshop_service import sync_workshops_master_and_links


class HardeningTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.password_hash = hash_password("Review123!")

    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        @event.listens_for(self.engine, "connect")
        def foreign_keys(conn, _):
            conn.execute("PRAGMA foreign_keys=ON")
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, autoflush=False)
        with self.sessions() as db:
            a, b = Workshop(name="哈东"), Workshop(name="哈南")
            db.add_all([a, b]); db.flush()
            admin = User(username="review_admin", role="super_admin", password_hash=self.password_hash)
            section_admin = User(username="review_section_admin", role="section_admin", password_hash=self.password_hash)
            manager = User(username="review_manager", role="workshop_director", workshop_id=a.id, password_hash=self.password_hash)
            workshop_admin = User(username="review_workshop_admin", role="workshop_admin", workshop_id=a.id, password_hash=self.password_hash)
            driver_user = User(username="review_driver", role="vehicle_driver", workshop_id=a.id, password_hash=self.password_hash)
            db.add_all([admin, section_admin, manager, workshop_admin, driver_user]); db.flush()
            ca = Vehicle(plate_number="TEST-A", workshop_id=a.id, org_unit=a.name)
            cb = Vehicle(plate_number="TEST-B", workshop_id=b.id, org_unit=b.name)
            driver = Driver(name="Synthetic", phone="00000000000", id_card="synthetic-id", workshop_id=a.id, user_id=driver_user.id)
            db.add_all([ca, cb, driver]); db.flush()
            trip = TripRequest(purpose="Review", applicant_name="Test", start_at=datetime(2026, 9, 1, 8), end_at=datetime(2026, 9, 1, 9),
                               origin="A", destination="B", created_by=admin.id, vehicle_id=ca.id, driver_id=driver.id, workshop_id=a.id, status="approved")
            db.add_all([trip, NavPreset(name="A", lng=126, lat=45), NavPreset(name="B", lng=126, lat=45)])
            db.commit()
            self.ids = dict(admin=admin.id, section_admin=section_admin.id, manager=manager.id,
                            workshop_admin=workshop_admin.id, user=driver_user.id,
                            a=a.id, b=b.id, ca=ca.id, cb=cb.id, driver=driver.id, trip=trip.id)
        self.actor = self.ids["admin"]
        def db_override():
            with self.sessions() as db:
                yield db
        def user_override():
            with self.sessions() as db:
                return db.get(User, self.actor)
        app.dependency_overrides[get_db] = db_override
        app.dependency_overrides[get_current_user] = user_override
        self.client = TestClient(app, raise_server_exceptions=True)

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        self.engine.dispose()

    def test_workshop_scope_and_mutations(self):
        with self.sessions() as db:
            other_driver = Driver(name="Other", phone="00000000001", id_card="other-secret", workshop_id=self.ids["b"])
            db.add(other_driver); db.commit()
            other_driver_id = other_driver.id
        self.actor = self.ids["manager"]
        other = self.client.get(f"/api/v1/drivers/{other_driver_id}")
        self.assertEqual(other.status_code, 200)
        self.assertTrue(other.json()["is_restricted"])
        self.assertEqual(other.json()["phone"], "00000000001")
        self.assertIsNone(other.json()["id_card"])
        self.assertEqual([v["id"] for v in self.client.get("/api/v1/vehicles").json()], [self.ids["ca"], self.ids["cb"]])
        r = self.client.patch(f'/api/v1/vehicles/{self.ids["cb"]}', json={"remarks": "forbidden"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(self.client.patch(f'/api/v1/vehicles/{self.ids["ca"]}', json={"workshop_id": self.ids["b"]}).status_code, 403)
        self.assertEqual(self.client.patch(f'/api/v1/drivers/{self.ids["driver"]}', json={"status": "本单位"}).status_code, 200)
        self.assertEqual(self.client.patch(f'/api/v1/drivers/{other_driver_id}', json={"status": "本单位"}).status_code, 403)
        self.assertEqual(self.client.get("/api/v1/dashboard/summary").json()["vehicles_total"], 2)
        for path in ["/drivers/stats", "/drivers/filters", "/fuel/balances", "/fuel/records", "/repair-records/settlements/summary", "/reports/export/summary.csv"]:
            self.assertEqual(self.client.get("/api/v1" + path).status_code, 200, path)

    def test_missing_workshop_does_not_grant_global_scope(self):
        with self.sessions() as db:
            db.get(User, self.ids["manager"]).workshop_id = None; db.commit()
        self.actor = self.ids["manager"]
        self.assertEqual(len(self.client.get("/api/v1/vehicles").json()), 2)
        self.assertEqual(self.client.patch(f'/api/v1/vehicles/{self.ids["ca"]}', json={"remarks": "blocked"}).status_code, 403)
        driver = self.client.get(f'/api/v1/drivers/{self.ids["driver"]}').json()
        self.assertTrue(driver["is_restricted"])

    def test_fuel_delete_respects_workshop_scope(self):
        with self.sessions() as db:
            row = FuelRecord(card_asn="OTHER", occur_time=datetime(2026, 9, 1), amount=10,
                             balance=20, volumn=2, workshop_id=self.ids["b"])
            db.add(row); db.commit()
            record_id = row.id
        self.actor = self.ids["manager"]
        self.assertEqual(self.client.delete(f"/api/v1/fuel/records/{record_id}").status_code, 403)
        with self.sessions() as db:
            self.assertIsNotNone(db.get(FuelRecord, record_id))
        self.actor = self.ids["admin"]
        self.assertEqual(self.client.delete(f"/api/v1/fuel/records/{record_id}").status_code, 204)

    def test_driver_sensitive_fields_are_hidden(self):
        self.actor = self.ids["user"]
        item = self.client.get("/api/v1/drivers").json()["items"][0]
        self.assertIsNone(item["id_card"])
        self.assertIsNone(item["health_check_report"])

    def test_driver_document_upload_download_and_delete(self):
        old_backend = settings.media_storage_backend
        old_upload_dir = settings.upload_dir
        with tempfile.TemporaryDirectory() as upload_dir:
            settings.media_storage_backend = "local"
            settings.upload_dir = upload_dir
            try:
                body = b"\x89PNG\r\n\x1a\n" + b"driver-health-report"
                path = f'/api/v1/drivers/{self.ids["driver"]}/documents/health_check_report'
                uploaded = self.client.post(path, files={"file": ("health.png", body, "image/png")})
                self.assertEqual(uploaded.status_code, 201)
                self.assertEqual(uploaded.json()["kind"], "health_check_report")

                driver = self.client.get(f'/api/v1/drivers/{self.ids["driver"]}').json()
                self.assertTrue(driver["has_health_check_report"])
                self.assertIsNone(driver["health_check_report"])
                downloaded = self.client.get(path)
                self.assertEqual(downloaded.status_code, 200)
                self.assertEqual(downloaded.content, body)
                replacement = b"\x89PNG\r\n\x1a\n" + b"replacement-health-report"
                replaced = self.client.post(path, files={"file": ("health-new.png", replacement, "image/png")})
                self.assertEqual(replaced.status_code, 201)
                self.assertEqual(self.client.get(path).content, replacement)
                patched = self.client.patch(
                    f'/api/v1/drivers/{self.ids["driver"]}',
                    json={"health_check_report": "media://local/_/forged.png"},
                )
                self.assertEqual(patched.status_code, 200)
                self.assertEqual(self.client.get(path).content, replacement)
                self.actor = self.ids["user"]
                self.assertEqual(self.client.get(path).status_code, 403)
                self.actor = self.ids["admin"]
                with self.sessions() as db:
                    self.assertEqual(db.scalar(select(func.count()).select_from(MediaAsset)), 1)

                self.assertEqual(self.client.delete(path).status_code, 204)
                self.assertEqual(self.client.get(path).status_code, 404)
                with self.sessions() as db:
                    self.assertEqual(db.scalar(select(func.count()).select_from(MediaAsset)), 0)
            finally:
                settings.media_storage_backend = old_backend
                settings.upload_dir = old_upload_dir

    def test_driver_account_binding_validates_role_and_scope(self):
        path = f'/api/v1/drivers/{self.ids["driver"]}'
        self.assertEqual(self.client.patch(path, json={"user_id": self.ids["admin"]}).status_code, 400)
        self.assertEqual(self.client.patch(path, json={"workshop_id": self.ids["b"]}).status_code, 400)
        self.actor = self.ids["manager"]
        with self.sessions() as db:
            other = User(username="other_driver", role="vehicle_driver", workshop_id=self.ids["b"], password_hash=self.password_hash)
            db.add(other); db.commit()
            other_id = other.id
        self.assertEqual(self.client.patch(path, json={"user_id": other_id}).status_code, 403)
        with self.sessions() as db:
            self.assertEqual(db.get(Driver, self.ids["driver"]).user_id, self.ids["user"])

    def test_director_cannot_delete_map_points(self):
        self.actor = self.ids["manager"]
        r = self.client.put("/api/v1/nav-presets", json={"markers": [{"name": "A", "lng": 126, "lat": 45}]})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(len(self.client.get("/api/v1/nav-presets").json()["markers"]), 2)

    def test_map_rejects_stale_save(self):
        initial = self.client.get("/api/v1/nav-presets").json()
        update = {"markers": [{"name": "A", "lng": 127, "lat": 45}], "revision": initial["revision"]}
        self.assertEqual(self.client.put("/api/v1/nav-presets", json=update).status_code, 200)
        self.assertEqual(self.client.put("/api/v1/nav-presets", json=initial).status_code, 409)
        self.assertEqual(self.client.get("/api/v1/nav-presets").json()["markers"][0]["lng"], 127)

    def test_driver_can_edit_own_pending_application(self):
        self.actor = self.ids["user"]
        r = self.client.post("/api/v1/trip-requests", json=dict(purpose="own", applicant_name="Test", origin="A", destination="B",
            start_at="2026-09-02T08:00:00Z", end_at="2026-09-02T09:00:00Z"))
        self.assertEqual(r.status_code, 201)
        url = f'/api/v1/trip-requests/{r.json()["id"]}'
        self.assertEqual(self.client.patch(url, json={"purpose": "updated", "notes": "ok"}).status_code, 200)
        self.assertEqual(self.client.patch(url, json={"status": "approved"}).status_code, 403)

    def test_account_generation_links_driver_and_is_repeatable(self):
        with self.sessions() as db:
            driver = Driver(name="Example", phone="00000000001", id_card="123456", workshop_id=self.ids["a"])
            db.add(driver); db.commit(); did = driver.id
        r = self.client.post("/api/v1/users/generate-driver-accounts")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["created"], 1)
        with self.sessions() as db:
            row = db.get(Driver, did)
            self.assertIsNotNone(row.user_id)
            self.assertEqual(db.get(User, row.user_id).workshop_id, self.ids["a"])
        self.assertEqual(self.client.post("/api/v1/users/generate-driver-accounts").json()["created"], 0)

    def test_repair_summary_paginates_and_preserves_totals(self):
        with self.sessions() as db:
            for i in range(3):
                rec = RepairRecord(vehicle_id=self.ids["ca"], created_by=self.ids["user"], repair_order_no=str(i))
                db.add(rec); db.flush()
                st = RepairSettlement(repair_record_id=rec.id, recognition_status="done", total_amount=10)
                db.add(st); db.flush()
                db.add(RepairSettlementLine(settlement_id=st.id, item_name="test", category_l1="parts", amount=10))
            db.commit()
        result = self.client.get("/api/v1/repair-records/settlements/summary?skip=1&limit=1").json()
        self.assertEqual(result["total"], 3)
        self.assertEqual(len(result["items"]), 1)
        self.assertEqual(float(result["category_totals"]["parts"]), 30)

    def test_all_account_levels_can_start_platform_sync(self):
        result = {"ok": True, "platform": "kunlun", "balance_written": 2, "record_written": 0}
        account_ids = ["admin", "section_admin", "manager", "workshop_admin", "user"]
        with patch("app.services.kunlun_balance_service.sync_kunlun_balances", return_value=result) as sync, \
             patch("app.services.fuel_sync_stats_service.record_fuel_sync_log"):
            for account_id in account_ids:
                with self.subTest(account_id=account_id):
                    self.actor = self.ids[account_id]
                    response = self.client.post("/api/v1/fuel/sync", json={})
                    self.assertEqual(response.status_code, 200)
                    self.assertIn('"type": "done"', response.text)
        self.assertEqual(sync.call_count, len(account_ids))

    def test_dashboard_vehicle_and_driver_totals_are_global_for_all_accounts(self):
        with self.sessions() as db:
            db.add(Driver(name="Other workshop driver", phone="00000000002", workshop_id=self.ids["b"]))
            db.commit()

        for account_id in ["admin", "section_admin", "manager", "workshop_admin", "user"]:
            with self.subTest(account_id=account_id):
                self.actor = self.ids[account_id]
                response = self.client.get("/api/v1/dashboard/summary")
                self.assertEqual(response.status_code, 200)
                payload = response.json()
                self.assertEqual(payload["vehicles_total"], 2)
                self.assertEqual(payload["drivers_total"], 2)

    def test_fuel_sync_stats_returns_explicit_utc_timestamp(self):
        with self.sessions() as db:
            db.add(FuelSyncLog(user_id=self.ids["admin"], synced_at=datetime(2026, 9, 24, 2, 25, 2), success=True))
            db.commit()

        response = self.client.get("/api/v1/fuel/sync-stats")
        self.assertEqual(response.status_code, 200)
        raw = response.json()["last_synced_at"]
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        self.assertEqual(parsed.utcoffset(), timedelta(0))
        self.assertEqual(parsed.hour, 2)

    def test_fuel_balances_are_global_for_all_accounts_and_exclude_retained_cards(self):
        with self.sessions() as db:
            db.add_all([
                FuelBalance(card_no="VISIBLE-A", workshop="哈东", vehicle_no="TEST-A", workshop_id=self.ids["a"], amount=100, reserve_fund=20, total=120),
                FuelBalance(card_no="VISIBLE-B", workshop="哈南", vehicle_no="TEST-B", workshop_id=self.ids["b"], amount=200, reserve_fund=30, total=230),
                FuelBalance(card_no="RETAINED-VEHICLE", workshop="/", vehicle_no=" 留存 ", amount=900, reserve_fund=100, total=1000),
                FuelBalance(card_no="RETAINED-WORKSHOP", workshop=" 留存 ", vehicle_no="KEEP", amount=900, reserve_fund=100, total=1000),
            ])
            db.commit()

        for account_id in ["admin", "section_admin", "manager", "workshop_admin", "user"]:
            with self.subTest(account_id=account_id):
                self.actor = self.ids[account_id]
                response = self.client.get("/api/v1/fuel/balances?sort_by=card_no&sort_dir=asc")
                self.assertEqual(response.status_code, 200)
                payload = response.json()
                self.assertEqual([item["card_no"] for item in payload["items"]], ["VISIBLE-A", "VISIBLE-B"])
                self.assertEqual(payload["total"], 2)
                self.assertEqual(float(payload["total_amount"]), 350)
                self.assertEqual(sum(bucket["count"] for bucket in payload["buckets"]), 2)
                self.assertEqual(self.client.get("/api/v1/fuel/balance-vehicles").json(), ["TEST-A", "TEST-B"])

        filtered = self.client.get("/api/v1/fuel/balances?vehicle_no=TEST-A").json()
        self.assertEqual([item["card_no"] for item in filtered["items"]], ["VISIBLE-A"])

    def test_platform_sync_contract_uses_kunlun_balance_service(self):
        result = {"ok": True, "platform": "kunlun", "balance_written": 2, "record_written": 0}
        with patch("app.services.kunlun_balance_service.sync_kunlun_balances", return_value=result) as sync, \
             patch("app.services.fuel_sync_stats_service.record_fuel_sync_log"):
            response = self.client.post("/api/v1/fuel/sync", json={})
        self.assertEqual(response.status_code, 200)
        self.assertIn('"type": "done"', response.text)
        self.assertIn('"platform": "kunlun"', response.text)
        sync.assert_called_once()

    def test_bill_sync_dispatches_dates_without_balance_stats(self):
        result = {"ok": True, "record_written": 41, "record_inserted": 0, "record_updated": 41}
        with patch("app.services.kunlun_bill_service.sync_kunlun_bills", return_value=result) as bills, \
             patch("app.services.kunlun_balance_service.sync_kunlun_balances") as balances, \
             patch("app.services.fuel_sync_stats_service.record_fuel_sync_log") as stats:
            response = self.client.post('/api/v1/fuel/sync', json={
                'target': 'bills', 'date_from': '2026-10-01', 'date_to': '2026-10-08'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('"record_updated": 41', response.text)
        self.assertEqual([str(value) for value in bills.call_args.args], ['2026-10-01', '2026-10-08'])
        balances.assert_not_called()
        stats.assert_not_called()

    def test_missing_bill_volume_is_blank_in_export(self):
        from openpyxl import load_workbook
        from io import BytesIO
        with self.sessions() as db:
            db.add(FuelRecord(card_asn='C', occur_time=datetime(2026, 10, 1), amount=10,
                              volumn=0, balance=0, volume_available=False, balance_available=False))
            db.commit()
        response = self.client.get('/api/v1/fuel/records/export.xlsx')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertIsNone(sheet.cell(2, 6).value)
        self.assertIsNone(sheet.cell(2, 7).value)
        self.assertIsNone(sheet.cell(2, 10).value)

    def test_bill_vehicle_filter_sort_and_export(self):
        from io import BytesIO
        from openpyxl import load_workbook
        with self.sessions() as db:
            db.add_all([
                FuelRecord(card_asn='C1', car_no='TEST-A', workshop_id=self.ids['a'], occur_time=datetime(2026, 10, 2), amount=20),
                FuelRecord(card_asn='C2', car_no='TEST-A', workshop_id=self.ids['a'], occur_time=datetime(2026, 10, 1), amount=10),
                FuelRecord(card_asn='C3', car_no='TEST-B', workshop_id=self.ids['b'], occur_time=datetime(2026, 10, 3), amount=30),
            ])
            db.commit()
        params = {'car_no': 'TEST-A', 'sort_by': 'occur_time', 'sort_dir': 'desc'}
        response = self.client.get('/api/v1/fuel/records', params=params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['total'], 2)
        self.assertEqual([float(r['amount']) for r in response.json()['items']], [20, 10])
        sheet = load_workbook(BytesIO(self.client.get('/api/v1/fuel/records/export.xlsx', params=params).content)).active
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual([float(sheet.cell(i, 8).value) for i in (2, 3)], [20, 10])
        by_vehicle = self.client.get('/api/v1/fuel/records', params={'sort_by': 'car_no', 'sort_dir': 'desc'}).json()
        self.assertEqual([r['car_no'] for r in by_vehicle['items']], ['TEST-B', 'TEST-A', 'TEST-A'])

    def test_bill_vehicle_options_respect_workshop_scope(self):
        with self.sessions() as db:
            db.add_all([
                FuelRecord(card_asn='C-A', car_no='TEST-A', workshop_id=self.ids['a'], occur_time=datetime(2026, 10, 1)),
                FuelRecord(card_asn='C-B', car_no='TEST-B', workshop_id=self.ids['b'], occur_time=datetime(2026, 10, 1)),
            ])
            db.commit()
        self.actor = self.ids['manager']
        self.assertEqual(self.client.get('/api/v1/fuel/record-vehicles').json(), ['TEST-A'])
        self.assertEqual(self.client.get('/api/v1/fuel/records', params={'car_no': 'TEST-B'}).json()['total'], 0)

    def test_bill_detail_scope_and_product_cache(self):
        with self.sessions.begin() as db:
            row = FuelRecord(card_asn='C-B', workshop_id=self.ids['b'], occur_time=datetime(2026, 10, 1),
                             platform_data={'orderNo': 'ORDER', 'driverName': 'Synthetic driver'})
            db.add(row); db.flush()
            record_id = row.id
        with patch('app.services.kunlun_bill_service.fetch_bill_products', return_value=[{'productName': 'Diesel', 'unit': 'L'}]) as fetch:
            self.actor = self.ids['manager']
            self.assertEqual(self.client.get(f'/api/v1/fuel/records/{record_id}').status_code, 404)
            fetch.assert_not_called()
            self.actor = self.ids['admin']
            detail = self.client.get(f'/api/v1/fuel/records/{record_id}')
            self.assertEqual(detail.status_code, 200)
            self.assertEqual(detail.json()['platform_data']['productDetails'][0]['productName'], 'Diesel')
            self.assertEqual(self.client.get(f'/api/v1/fuel/records/{record_id}').status_code, 200)
            fetch.assert_called_once_with('ORDER')
            listing = self.client.get('/api/v1/fuel/records').json()['items'][0]
            self.assertNotIn('platform_data', listing)

    def test_bill_detail_platform_failure_keeps_bill_and_allows_retry(self):
        with self.sessions.begin() as db:
            row = FuelRecord(card_asn='C-A', workshop_id=self.ids['a'], occur_time=datetime(2026, 10, 1),
                             platform_data={'orderNo': 'ORDER'}, amount=100)
            db.add(row); db.flush()
            record_id = row.id
        with patch('app.services.kunlun_bill_service.fetch_bill_products', side_effect=RuntimeError('synthetic failure')):
            result = self.client.get(f'/api/v1/fuel/records/{record_id}')
            self.assertEqual(result.status_code, 200)
            self.assertEqual(float(result.json()['amount']), 100)
            self.assertIsNotNone(result.json()['product_detail_error'])
            self.assertNotIn('productDetails', result.json()['platform_data'])
        with patch('app.services.kunlun_bill_service.fetch_bill_products', return_value=[]):
            result = self.client.get(f'/api/v1/fuel/records/{record_id}').json()
            self.assertIsNone(result['product_detail_error'])
            self.assertEqual(result['platform_data']['productDetails'], [])

    def test_bill_missing_card_filter_keeps_count_and_export_consistent(self):
        from io import BytesIO
        from openpyxl import load_workbook
        with self.sessions() as db:
            db.add_all([FuelRecord(card_asn=card, car_no=plate, occur_time=datetime(2026, 10, 1))
                        for card, plate in [('CARD', 'VALID'), ('', 'EMPTY'), ('   ', 'SPACE')]])
            db.commit()
        params = {'has_card_only': True, 'page_size': 1}
        response = self.client.get('/api/v1/fuel/records', params=params).json()
        self.assertEqual(response['total'], 1)
        self.assertEqual([r['car_no'] for r in response['items']], ['VALID'])
        sheet = load_workbook(BytesIO(self.client.get('/api/v1/fuel/records/export.xlsx', params=params).content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(self.client.get('/api/v1/fuel/record-vehicles').json(), ['VALID'])

    def test_fuel_filter_uses_shanghai_calendar_days(self):
        with self.sessions() as db:
            db.add(FuelRecord(card_asn="C", occur_time=datetime(2026, 9, 1, 17), amount=10, balance=20, volumn=2))
            db.commit()
        self.assertEqual(self.client.get("/api/v1/fuel/records?date_from=2026-09-02&date_to=2026-09-02").json()["total"], 1)

    def test_driver_trip_transition_and_reopen_rejected(self):
        self.actor = self.ids["user"]
        url = f'/api/v1/trip-requests/{self.ids["trip"]}'
        for state in ["in_progress", "completed"]:
            self.assertEqual(self.client.patch(url, json={"status": state, "notes": "test"}).status_code, 200)
        self.assertEqual(self.client.patch(url, json={"status": "in_progress"}).status_code, 409)

    def test_duplicate_reservation_is_rejected(self):
        r = self.client.post("/api/v1/trip-requests", json=dict(purpose="other", applicant_name="Test", origin="A", destination="B",
            start_at="2026-09-01T08:30:00Z", end_at="2026-09-01T09:30:00Z", vehicle_id=self.ids["ca"], driver_id=self.ids["driver"]))
        self.assertEqual(r.status_code, 201)
        r = self.client.patch(f'/api/v1/trip-requests/{r.json()["id"]}', json={"status": "approved"})
        self.assertEqual(r.status_code, 409)

    def test_null_update_validation(self):
        self.assertEqual(self.client.patch(f'/api/v1/trip-requests/{self.ids["trip"]}', json={"start_at": None}).status_code, 422)
        self.assertEqual(self.client.patch(f'/api/v1/vehicles/{self.ids["ca"]}', json={"plate_number": None}).status_code, 422)
        self.assertEqual(self.client.patch(f'/api/v1/vehicles/{self.ids["ca"]}', json={"remarks": None}).status_code, 200)

    def test_locked_sqlite_login_returns_423(self):
        with self.sessions() as db:
            db.get(User, self.ids["user"]).lock_until = datetime.now(timezone.utc) + timedelta(minutes=10); db.commit()
        r = self.client.post("/api/v1/auth/login", json={"username": "review_driver", "password": "Review123!"})
        self.assertEqual(r.status_code, 423)

    def test_startup_preserves_user_and_workshop(self):
        with self.sessions() as db:
            db.get(Workshop, self.ids["a"]).is_active = False; db.commit()
            sync_workshops_master_and_links(db)
            sync_workshops_master_and_links(db)
            self.assertEqual(db.get(User, self.ids["user"]).workshop_id, self.ids["a"])
            self.assertEqual(db.get(User, self.ids["manager"]).username, "review_manager")
            self.assertFalse(db.get(Workshop, self.ids["a"]).is_active)

    def test_failed_recognition_persists_and_preserves_lines(self):
        with self.sessions() as db:
            rec = RepairRecord(vehicle_id=self.ids["ca"], created_by=self.ids["user"], repair_order_no="R", photo_settlement_path=str(Path(__file__)))
            db.add(rec); db.flush()
            settlement = RepairSettlement(repair_record_id=rec.id, recognition_status="done")
            db.add(settlement); db.flush()
            db.add(RepairSettlementLine(settlement_id=settlement.id, item_name="keep", amount=10)); db.commit()
            with patch("app.services.settlement_recognition_service.recognize_settlement_image", side_effect=RuntimeError("synthetic")):
                result = recognize_and_save(db, rec); db.commit()
            self.assertEqual(result.recognition_status, "failed")
            self.assertTrue(result.recognition_error)
            self.assertEqual(db.scalar(select(func.count()).select_from(RepairSettlementLine)), 1)
            self.assertEqual(rec.status, "recognize_failed")

    def test_kunlun_balance_import_updates_current_balance_and_mapping(self):
        with self.sessions() as db:
            db.add(FuelCardLookup(card_no="CARD", workshop="哈东", vehicle_no="TEST-A"))
            db.add(FuelBalance(card_no="CARD", amount=100, reserve_fund=10, total=110))
            db.commit()
        rows = [{"卡号": "CARD", "当前余额": "80.25"}, {"卡号": "UNMATCHED", "当前余额": 20}]
        result = import_balance_rows(rows, session_factory=self.sessions)
        self.assertEqual(result["balance_written"], 2)
        self.assertEqual(result["balance_matched"], 1)
        with self.sessions() as db:
            linked = db.scalar(select(FuelBalance).where(FuelBalance.card_no == "CARD"))
            unlinked = db.scalar(select(FuelBalance).where(FuelBalance.card_no == "UNMATCHED"))
            self.assertEqual(linked.total, 80.25)
            self.assertEqual(linked.reserve_fund, 0)
            self.assertEqual(linked.workshop_id, self.ids["a"])
            self.assertEqual(unlinked.total, 20)
            self.assertEqual(unlinked.workshop, "")

    def test_kunlun_balance_import_is_atomic_on_invalid_payload(self):
        with self.sessions() as db:
            db.add(FuelBalance(card_no="CARD", amount=100, reserve_fund=0, total=100)); db.commit()
        with self.assertRaises(RuntimeError):
            import_balance_rows(
                [{"卡号": "CARD", "当前余额": 1}, {"卡号": "OTHER", "当前余额": "bad"}],
                session_factory=self.sessions,
            )
        with self.sessions() as db:
            self.assertEqual(db.scalar(select(FuelBalance.total)), 100)
            self.assertEqual(db.scalar(select(func.count()).select_from(FuelBalance)), 1)

    def test_fuel_analysis_does_not_invent_mileage(self):
        rows = [FuelRecord(car_no="A", card_asn="1", occur_time=datetime(2026, 1, 1), volumn=10, amount=50),
                FuelRecord(car_no="B", card_asn="2", occur_time=datetime(2026, 1, 1), volumn=20, amount=100)]
        analysis = build_fuel_analysis(rows, {})
        self.assertEqual({x.plate_number for x in analysis.vehicles}, {"A", "B"})
        self.assertIsNone(analysis.overview.avg_l_per_100km)

    def test_dashboard_aggregates_fuel_in_database_without_losing_record_count(self):
        with self.sessions() as db:
            db.add_all([
                FuelRecord(card_asn="CARD-A", car_no="TEST-A", occur_time=datetime(2026, 1, 1), volumn=10, amount=50),
                FuelRecord(card_asn="CARD-A", car_no="TEST-A", occur_time=datetime(2026, 1, 2), volumn=20, amount=100),
            ])
            db.commit()

        response = self.client.get("/api/v1/dashboard/summary")
        self.assertEqual(response.status_code, 200)
        analysis = response.json()["fuel_analysis"]
        self.assertEqual(analysis["overview"]["records"], 2)
        self.assertEqual(analysis["overview"]["total_liters"], 30.0)
        self.assertEqual(analysis["overview"]["total_amount"], 150.0)
        self.assertEqual(analysis["vehicles"][0]["vehicle_id"], self.ids["ca"])

    def test_additive_migration_keeps_legacy_data(self):
        from app.db import migrate
        with self.engine.begin() as conn:
            conn.execute(text("CREATE TABLE bus_expense (id INTEGER PRIMARY KEY, amount INTEGER)"))
            conn.execute(text("INSERT INTO bus_expense VALUES (1, 100)"))
        with patch.object(migrate, "engine", self.engine):
            migrate.run_runtime_migrations(); migrate.run_runtime_migrations()
        with self.engine.connect() as conn:
            self.assertEqual(conn.scalar(text("SELECT amount FROM bus_expense WHERE id=1")), 100)

    def test_legacy_table_rename_preserves_rows_and_foreign_keys(self):
        from app.db import migrate
        legacy_engine = create_engine("sqlite://")
        with legacy_engine.begin() as conn:
            conn.execute(text("PRAGMA foreign_keys=ON"))
            conn.execute(text("CREATE TABLE bus_workshop (id INTEGER PRIMARY KEY, name TEXT NOT NULL)"))
            conn.execute(
                text(
                    "CREATE TABLE bus_vehicle ("
                    "id INTEGER PRIMARY KEY, plate_number TEXT NOT NULL, "
                    "workshop_id INTEGER REFERENCES bus_workshop(id))"
                )
            )
            conn.execute(text("INSERT INTO bus_workshop VALUES (1, '哈东')"))
            conn.execute(text("INSERT INTO bus_vehicle VALUES (7, 'TEST-7', 1)"))
        with patch.object(migrate, "engine", legacy_engine):
            migrate.run_pre_create_migrations()
            migrate.run_pre_create_migrations()
        inspector = inspect(legacy_engine)
        self.assertNotIn("bus_vehicle", inspector.get_table_names())
        self.assertIn("vehicles", inspector.get_table_names())
        self.assertEqual(inspector.get_foreign_keys("vehicles")[0]["referred_table"], "workshops")
        with legacy_engine.connect() as conn:
            self.assertEqual(conn.scalar(text("SELECT plate_number FROM vehicles WHERE id=7")), "TEST-7")
        legacy_engine.dispose()

    def test_every_model_table_has_a_business_schema(self):
        from app.db.structure import TABLE_SCHEMAS

        self.assertEqual(set(Base.metadata.tables), set(TABLE_SCHEMAS))
        self.assertEqual(TABLE_SCHEMAS["fuel_cards"], "fuel_management")
        self.assertEqual(TABLE_SCHEMAS["vehicles"], "fleet_management")


if __name__ == "__main__":
    unittest.main()
