"""账单跨页筛选、多级排序和导出回归，使用独立内存库。"""
import os
os.environ.update(DATABASE_URL="sqlite://", ENVIRONMENT="test", DEMO_SEEDING_ENABLED="false")

import asyncio
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
import unittest

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.dialects import postgresql
from sqlalchemy.orm import Session

import app.models
from app.api.v1.fuel import _record_ordering, export_records_xlsx, list_records, list_record_card_options
from app.db.base import Base
from app.models.fuel import FuelRecord
from app.models.workshop import Workshop


class FuelRecordSortingTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.user = SimpleNamespace(role="super_admin", workshop_id=None)
        # UTC 16:00 跨过上海零点；同一天内车号优先于交易时分。
        self.db.add_all([
            FuelRecord(id=1, car_no="B", card_asn="CARD-B", workshop="W", occur_time=datetime(2026, 10, 7, 16, 30)),
            FuelRecord(id=2, car_no="A", card_asn="CARD-A", workshop="W", occur_time=datetime(2026, 10, 8, 2)),
            FuelRecord(id=3, car_no="C", card_asn="CARD-C", workshop="W", occur_time=datetime(2026, 10, 7, 15, 59)),
            FuelRecord(id=4, car_no="A", card_asn="CARD-A", workshop="W", occur_time=datetime(2026, 10, 8, 3)),
            FuelRecord(id=5, car_no="A", card_asn="", workshop="W", occur_time=datetime(2026, 10, 6, 1)),
        ])
        self.db.commit()

    def tearDown(self):
        self.db.close()
        self.engine.dispose()

    def page(self, **params):
        return list_records(self.db, self.user, has_card_only=True, sort_by="occur_time", sort_dir="asc",
                            sort_secondary_by="car_no", **params)

    def test_date_then_car_sort_is_global_across_pages(self):
        pages = [self.page(page=p, page_size=2) for p in (1, 2)]
        self.assertEqual([r.id for page in pages for r in page.items], [3, 2, 4, 1])
        self.assertEqual([p.total for p in pages], [4, 4])

    def test_secondary_descending_preserves_primary_date(self):
        self.assertEqual([r.id for r in self.page(sort_secondary_dir="desc").items], [3, 1, 2, 4])

    def test_car_primary_and_date_secondary(self):
        result = list_records(self.db, self.user, has_card_only=True, sort_by="car_no", sort_dir="asc",
                              sort_secondary_by="occur_time", sort_secondary_dir="desc")
        self.assertEqual([r.id for r in result.items], [4, 2, 1, 3])

    def test_filters_apply_before_pagination_and_shanghai_day_boundary(self):
        result = self.page(car_no="A", card_asn="CARD-A", date_from=date(2026, 10, 8),
                           date_to=date(2026, 10, 8), page_size=1, page=2)
        self.assertEqual(result.total, 2)
        self.assertEqual([r.id for r in result.items], [4])

    def test_single_field_sort_remains_timestamp_based(self):
        result = list_records(self.db, self.user, has_card_only=True, sort_by="occur_time", sort_dir="asc")
        self.assertEqual([r.id for r in result.items], [3, 1, 2, 4])

    def test_export_order_matches_all_list_pages(self):
        response = export_records_xlsx(self.db, self.user, has_card_only=True, sort_by="occur_time", sort_dir="asc",
                                       sort_secondary_by="car_no", sort_secondary_dir="asc")
        async def read():
            return b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(BytesIO(asyncio.run(read())))
        self.assertEqual([row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)], [3, 2, 4, 1])

    def test_postgres_sort_uses_shanghai_timezone(self):
        db = SimpleNamespace(get_bind=lambda: SimpleNamespace(dialect=SimpleNamespace(name="postgresql")))
        stmt = select(FuelRecord.id).order_by(*_record_ordering(db, "occur_time", "asc", "car_no", "asc"))
        sql = str(stmt.compile(dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True}))
        self.assertIn("date(timezone('Asia/Shanghai', fuel_records.occur_time)) ASC", sql)

    def test_card_options_use_latest_visible_vehicle_and_deduplicate_cards(self):
        self.db.add_all([Workshop(id=1, name="W"), Workshop(id=2, name="OTHER")])
        self.db.flush()
        self.db.get(FuelRecord, 2).car_no = "OLD"
        self.db.add(FuelRecord(id=6, car_no="PRIVATE", card_asn="CARD-A", workshop="OTHER", workshop_id=2,
                               occur_time=datetime(2026, 10, 9, 1)))
        self.db.commit()
        options = list_record_card_options(self.db, self.user)
        self.assertEqual(options, [
            {"card_asn": "CARD-A", "car_no": "PRIVATE"},
            {"card_asn": "CARD-B", "car_no": "B"},
            {"card_asn": "CARD-C", "car_no": "C"},
        ])
        visible = list_record_card_options(self.db, SimpleNamespace(role="workshop_admin", workshop_id=1))
        self.assertEqual(visible[0], {"card_asn": "CARD-A", "car_no": "A"})
        self.assertEqual(len(visible), 3)


if __name__ == "__main__":
    unittest.main()
