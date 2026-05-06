"""从 Excel 导入油卡余额快照到 bus_fuel_balance。

默认文件：D:\\Desktop\\获取油卡余额\\卡内剩余金额.xlsx
默认工作表：Sheet1
列约定：卡号、车间、车号、金额、备用金、合计
"""
from __future__ import annotations

import argparse
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.db.base import Base  # noqa: E402
from app.db.session import SessionLocal, engine  # noqa: E402
from app.models.fuel import FuelBalance  # noqa: E402

DEFAULT_XLSX = Path(r"D:\Desktop\获取油卡余额\卡内剩余金额.xlsx")
SHEET_NAME = "Sheet1"


def as_text(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def as_decimal(v) -> Decimal:
    s = as_text(v)
    if not s:
        return Decimal("0")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def run_import(path: Path, *, dry_run: bool) -> tuple[int, int]:
    if not path.is_file():
        raise SystemExit(f"文件不存在: {path}")

    Base.metadata.create_all(bind=engine)
    wb = load_workbook(path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"工作表不存在: {SHEET_NAME}")
    ws = wb[SHEET_NAME]

    header = [as_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    need = ["卡号", "车间", "车号", "金额", "备用金", "合计"]
    idx = {name: header.index(name) + 1 for name in need if name in header}
    missing = [x for x in need if x not in idx]
    if missing:
        raise SystemExit(f"缺少列: {missing}")

    db = SessionLocal()
    try:
        rows: list[dict[str, object]] = []
        for r in range(2, ws.max_row + 1):
            card_no = as_text(ws.cell(r, idx["卡号"]).value)
            if not card_no:
                continue
            rows.append(
                {
                    "card_no": card_no,
                    "workshop": as_text(ws.cell(r, idx["车间"]).value),
                    "vehicle_no": as_text(ws.cell(r, idx["车号"]).value),
                    "amount": as_decimal(ws.cell(r, idx["金额"]).value),
                    "reserve_fund": as_decimal(ws.cell(r, idx["备用金"]).value),
                    "total": as_decimal(ws.cell(r, idx["合计"]).value),
                }
            )

        created = 0
        updated = 0
        for row in rows:
            existed = db.scalar(select(FuelBalance).where(FuelBalance.card_no == row["card_no"]))
            if existed:
                existed.workshop = str(row["workshop"])
                existed.vehicle_no = str(row["vehicle_no"])
                existed.amount = row["amount"]  # type: ignore[assignment]
                existed.reserve_fund = row["reserve_fund"]  # type: ignore[assignment]
                existed.total = row["total"]  # type: ignore[assignment]
                updated += 1
            else:
                db.add(FuelBalance(**row))
                created += 1

        if dry_run:
            db.rollback()
        else:
            db.commit()
        return created, updated
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main() -> None:
    ap = argparse.ArgumentParser(description="导入油卡余额快照（Sheet1）到 bus_fuel_balance")
    ap.add_argument("--file", type=Path, default=DEFAULT_XLSX, help=f"Excel 路径，默认 {DEFAULT_XLSX}")
    ap.add_argument("--dry-run", action="store_true", help="仅解析并统计，不写库")
    args = ap.parse_args()

    created, updated = run_import(args.file, dry_run=args.dry_run)
    if args.dry_run:
        print(f"dry-run：预计新建 {created} 条，更新 {updated} 条")
    else:
        print(f"完成：bus_fuel_balance 新建 {created} 条，更新 {updated} 条")


if __name__ == "__main__":
    main()
