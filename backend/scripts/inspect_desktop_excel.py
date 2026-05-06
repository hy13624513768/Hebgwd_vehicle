"""临时脚本：查看桌面「数据」目录下 Excel 的表名与列名（UTF-8 输出）。"""
from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("需要: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(r"C:\Users\HY\Desktop\数据")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8") if hasattr(sys.stdout, "reconfigure") else None
    if not ROOT.is_dir():
        print(f"目录不存在: {ROOT}", file=sys.stderr)
        sys.exit(1)
    for path in sorted(ROOT.glob("*.xlsx")):
        print("=" * 60)
        print(path.name)
        wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
        for sn in wb.sheetnames[:5]:
            ws = wb[sn]
            print(f"  [sheet] {sn}")
            for r in range(1, 16):
                vals = []
                for c in range(1, 14):
                    v = ws.cell(r, c).value
                    vals.append(v)
                if any(v is not None for v in vals):
                    print(f"    R{r}: {vals}")
        wb.close()


if __name__ == "__main__":
    main()
